"""src/visualization/fonts.py ＋ src/output/docx_styles.py 契约测试（s918-31）。

字体侧只测回退链与缓存（monkeypatch 文件系统可见性，不碰真实 Windows\\Fonts）；
样式侧钉常量现值。任务书提到的「matplotlib 缓存」与磁盘不符——本模块不 import matplotlib，
缓存是模块级 `_FONT_FAMILY`，按现状测。
"""
import io
import os
import re
import zipfile
from types import SimpleNamespace

import pytest
from docx import Document
from docx.oxml.ns import qn

from src.output import docx_exporter
from src.output import docx_styles as ds
from src.visualization import fonts as fnt

ALL_CANDIDATES = [
    "msyh.ttc", "msyhbd.ttc", "simhei.ttf", "simsun.ttc",
    "simkai.ttf", "msjh.ttc", "Deng.ttf", "Dengb.ttf",
]


@pytest.fixture(autouse=True)
def reset_font_cache(monkeypatch):
    monkeypatch.setattr(fnt, "_FONT_FAMILY", None)


def _fake_fs(monkeypatch, tmp_path, existing=(), glob_hits=()):
    """把 WINDIR 指到 tmp，并只让 `existing` 里的文件名「存在」。"""
    monkeypatch.setenv("WINDIR", str(tmp_path))
    real_exists = os.path.exists

    def exists(path):
        name = os.path.basename(str(path))
        return name in existing and real_exists(str(tmp_path))

    monkeypatch.setattr(fnt.os.path, "exists", exists)
    monkeypatch.setattr(fnt, "glob", SimpleNamespace(glob=lambda pattern: [str(tmp_path / "Fonts" / h) for h in glob_hits]))
    return tmp_path / "Fonts"


def test_priority_order_prefers_yahei_over_simhei(monkeypatch, tmp_path):
    root = _fake_fs(monkeypatch, tmp_path, existing={"msyh.ttc", "simhei.ttf"})

    assert fnt.find_chinese_font() == os.path.join(str(root), "msyh.ttc")


def test_falls_to_next_candidate_when_yahei_absent(monkeypatch, tmp_path):
    root = _fake_fs(monkeypatch, tmp_path, existing={"simhei.ttf", "simsun.ttc"})

    assert fnt.find_chinese_font() == os.path.join(str(root), "simhei.ttf")


def test_every_candidate_is_tried_in_declared_order(monkeypatch, tmp_path):
    _fake_fs(monkeypatch, tmp_path, existing=set())
    seen = []
    monkeypatch.setattr(fnt.os.path, "exists", lambda p: seen.append(os.path.basename(str(p))) or False)

    assert fnt.find_chinese_font() is None
    assert seen[: len(ALL_CANDIDATES)] == ALL_CANDIDATES


def test_wildcard_search_is_the_last_resort(monkeypatch, tmp_path):
    root = _fake_fs(monkeypatch, tmp_path, existing=set(), glob_hits=["msyh.ttc"])

    assert fnt.find_chinese_font() == os.path.join(str(root), "msyh.ttc")


def test_returns_none_when_nothing_found(monkeypatch, tmp_path):
    """旧名后缀「though_annotated_str」记的是注解与实现不符；s919-07 已把注解改成 Optional[str]，后缀失去意义。"""
    _fake_fs(monkeypatch, tmp_path, existing=set(), glob_hits=[])

    assert fnt.find_chinese_font() is None
    assert fnt.get_font_family() == "sans-serif"


@pytest.mark.parametrize(
    "filename,family",
    [
        ("msyh.ttc", "Microsoft YaHei"),
        ("msyhbd.ttc", "Microsoft YaHei"),
        ("simhei.ttf", "SimHei"),
        ("simsun.ttc", "SimSun"),
        ("simkai.ttf", "KaiTi"),
        ("Deng.ttf", "DengXian"),                  # s919-07 翻正（原钉 sans-serif＝方框现状）
        ("Dengb.ttf", "DengXian"),                 # 等线粗体同名族
        ("msjh.ttc", "Microsoft JhengHei"),        # s919-07 翻正（微软正黑体）
    ],
)
def test_family_name_mapping_covers_six_families(monkeypatch, tmp_path, filename, family):
    """族名字符串取自本机 Fonts 目录的 name 表 ID=1（fontTools 实测），不是猜的别名。"""
    root = _fake_fs(monkeypatch, tmp_path, existing={filename})

    assert fnt.get_font_family() == family


def test_only_unknown_font_falls_back_to_sans_serif(monkeypatch, tmp_path):
    """六个族名之外的字体仍须落回 sans-serif（翻正不能顺手放宽兜底）。"""
    _fake_fs(monkeypatch, tmp_path, existing={"arial.ttf"})

    assert fnt.get_font_family() == "sans-serif"


def test_find_chinese_font_annotation_is_optional():
    """`find_chinese_font()` 一直可能返回 None，注解 s919-07 起改为 Optional[str]。

    钉注解而非行为（行为由 `test_returns_none_when_nothing_found` 覆盖）：
    注解与实现不符会让调用方按 str 做字符串拼接而静默炸。
    """
    import typing

    ret = typing.get_type_hints(fnt.find_chinese_font)["return"]
    args = typing.get_args(ret)
    assert type(None) in args, f"返回注解应允许 None，实得 {ret}"


def test_chinese_font_name_is_cached_after_first_call(monkeypatch, tmp_path):
    _fake_fs(monkeypatch, tmp_path, existing={"simhei.ttf"})
    calls = []
    monkeypatch.setattr(fnt, "get_font_family", lambda: (calls.append(1), "SimHei")[1])

    assert fnt.get_chinese_font() == "SimHei"
    assert fnt.get_chinese_font() == "SimHei"
    assert len(calls) == 1


def test_family_mapping_lowercases_the_discovered_path(monkeypatch, tmp_path):
    """glob 可带回大写文件名（Windows 实际会），族名判定走 .lower()。"""
    _fake_fs(monkeypatch, tmp_path, existing=set(), glob_hits=["MSYH.TTC"])

    assert fnt.get_font_family() == "Microsoft YaHei"


def test_docx_latin_and_cjk_font_constants_pinned():
    assert ds.FONT_BODY_LATIN == "Times New Roman"
    assert ds.FONT_BODY_CJK == "宋体"
    assert ds.FONT_HEADING_LATIN == "Times New Roman"
    assert ds.FONT_HEADING_CJK == "黑体"


def test_docx_size_and_spacing_constants_pinned():
    assert (ds.SIZE_TITLE, ds.SIZE_H1, ds.SIZE_H2, ds.SIZE_H3) == (16, 14, 13, 12)
    assert (ds.SIZE_BODY, ds.SIZE_CAPTION, ds.SIZE_TABLE) == (12, 10.5, 10.5)
    assert ds.LINE_SPACING == 1.5
    # s924-24：原为一条三元组断言，拆成三条独立断言以便单点定位（s923-18 三案共同前置）。
    # 期望值零变化，仍是 12 / 6。
    assert ds.SPACE_BEFORE_HEADING_PT == 12, (
        f"标题段前距常量现值漂移：期望 12，实得 {ds.SPACE_BEFORE_HEADING_PT}"
    )
    assert ds.SPACE_AFTER_HEADING_PT == 6, (
        f"标题段后距常量现值漂移：期望 6，实得 {ds.SPACE_AFTER_HEADING_PT}"
    )
    # s928-01 第一笔（案 B′ 改值）：期望值 0 → 10，与 `docx_styles.py:24` **同批**入库。
    # 本例只钉常量现值；产物面那一格看下面 `test_docx_body_paragraph_after_spacing_*`。
    assert ds.SPACE_AFTER_BODY_PT == 10, (
        f"正文段后距常量现值漂移：期望 10（B′ 本仓锁定），实得 {ds.SPACE_AFTER_BODY_PT}"
    )
    assert ds.FIRST_LINE_INDENT_CM == 0.74


def test_docx_apa7_margins_and_table_style_pinned():
    assert (ds.MARGIN_TOP_CM, ds.MARGIN_BOTTOM_CM) == (2.54, 2.54)
    assert (ds.MARGIN_LEFT_CM, ds.MARGIN_RIGHT_CM) == (3.18, 3.18)
    assert ds.TABLE_HEADER_FILL == "DDDDDD"
    assert ds.TABLE_BORDER_COLOR == "000000"
    assert ds.TABLE_BORDER_SIZE == 4
    assert (ds.FIG_PREFIX, ds.TABLE_PREFIX) == ("图", "表")


# ── 正文段后距的产物面看守（s924-25 建例，s928-01 两笔重评为 B′ 态）───────────────
# 建例时钉的是「常量说 0 / 产物给 10」这道**脱钩缝**（缝宽＝产物继承值 − 常量现值）。
# B′（s927-51 拍板、用户 2026-09-28 授权）＝ 第一笔把常量改 10 ＋ 第二笔接线 ⇒ 脱钩消失，
# 本例因此改钉三件**互相独立**的事实（而不是把 `10.0` 划成 `0.0`——那会退化成恒真式）：
# ①（臂一）正文段带**显式** `w:after`，值为整数字面量 "200"（＝10pt，本仓锁定值）；
# ②（臂二）产物 `docDefaults` 的 after 仍是 200 twips ⇒ direct 与继承同格 ⇒ **零版式变化**；
#    模板一旦变了（docDefaults ≠ 200）本臂即红——正是台账 241ⓒ 记的「未来不等价」那一条；
# ③（臂三）direct 读数与常量现值同格（容差＝半网格 0.025pt）⇒ 接线的因果链成立。
# 档二口径取自 docs/insights/s927-45-format-caliber-note.md §4-B：钉**整数字面量与属性名**，
# 不得用 `Pt(常量)` 派生期望值（否则退化成「两侧同过一个换算函数」的恒真式）。
DOCX_BODY_AFTER_EFFECTIVE_TWIPS = "200"   # 正文段 direct 与模板 docDefaults 共同的 w:after（twips 字面量）

_SYNTH_BODY = "合成正文段落（人造文案，不含任何研究或统计结论）"
_SYNTH_HEADING = "合成一级标题"


def _export_body_sample():
    """跑一次真实导出，返回 (正文段的 w:spacing 元素, 一级标题段的 w:spacing 元素, styles.xml)。

    走 `docx_exporter.build_review_docx`（纯代码生成，无模板）——它是 `docx_styles`
    唯一的实际接线方，正文段落的产路是 `_add_paragraph_with_inline_format`。
    """
    payload = docx_exporter.build_review_docx(
        title="合成综述",
        markdown_body=f"# {_SYNTH_HEADING}\n\n{_SYNTH_BODY}",
    )
    doc = Document(io.BytesIO(payload))
    spacing_of = {}
    for p in doc.paragraphs:
        ppr = p._p.find(qn("w:pPr"))
        sp = None if ppr is None else ppr.find(qn("w:spacing"))
        spacing_of[p.text.strip()] = sp
    with zipfile.ZipFile(io.BytesIO(payload)) as zf:
        styles_xml = zf.read("word/styles.xml").decode("utf-8")
    return spacing_of[_SYNTH_BODY], spacing_of[_SYNTH_HEADING], styles_xml


def _docdefaults_after_twips(styles_xml: str) -> int:
    """从产物包里读 docDefaults 的实际段后距（twips）——期望值不写死、跟模板走。"""
    m = re.search(r"<w:docDefaults>.*?<w:spacing[^>]*w:after=\"(\d+)\"", styles_xml, re.S)
    assert m, "产物 styles.xml 里找不到 docDefaults 的 w:after，本例前提失效"
    return int(m.group(1))


def test_docx_body_paragraph_after_spacing_effective_value_pinned():
    """正文段后距**产物生效值**快照（s928-01 两笔后的 B′ 态）。

    建例原委（s924-25）：导出器只在标题分支（`level > 0`）写 `pf.space_after`，正文分支
    从头到尾没碰过段后距 ⇒ 正文生效值由 `word/styles.xml` 的 `docDefaults after=200`（twips）
    继承而来是 **10pt**，而 `docx_styles.py` 里那个常量当时写 `0`、全仓消费方 **0 处**。
    当时两侧确实脱钩，所以 s924-25 把期望式写成「缝宽＝继承值 − 常量」。

    B′（s927-51 拍板、用户 2026-09-28 授权）落地后本例整例重评（卷宗口径：
    「谁改本常量必须连本断言一起评审」），且**不能**只把 `10.0` 划成 `0.0`——那样它就退化成
    s927-45 §4-B 禁的「两侧同过一个换算函数」的恒真式。三臂各自独立、可单点定位：
    ①（臂一）正文段现在**必须**带显式 `w:after == "200"`（整数字面量＝本仓锁定 10pt 已接线）；
    ②（臂二）模板 `docDefaults` 仍是 200 twips ⇒ direct 与继承同格 ⇒ **零版式变化**；
       模板变更时本臂红，正是「本仓锁定 vs 跟随模板」分叉的哨兵（台账 241ⓒ）；
    ③（臂三）direct 读数与常量现值同格（容差半网格 0.025pt）⇒ 接线的因果链成立。

    现状快照非规范：本例不主张 10pt 比 0pt 好，只主张「产物生效值可查、三处同格、来源明白」。
    """
    body_sp, heading_sp, styles_xml = _export_body_sample()

    # 先证明"探针不是恒真"：同一份产物里，接了线的标题段确实带 w:after，
    # 正文段的读数与之**分属两个元素**——二者对比才说明正文的读数是接线来的而不是读不到属性。
    assert heading_sp is not None, "一级标题段没有 w:spacing，导出器版式分支与本例前提不符"
    assert heading_sp.get(qn("w:after")) == str(int(ds.SPACE_AFTER_HEADING_PT * 20)), (
        f"标题段后距读数异常：期望 {ds.SPACE_AFTER_HEADING_PT}pt（{ds.SPACE_AFTER_HEADING_PT * 20:.0f} twips），"
        f"实得 w:after={heading_sp.get(qn('w:after'))!r}"
    )

    assert body_sp is not None, "正文段落连 w:pPr/w:spacing 都没有，产物版式与本例前提不符"
    # 臂一·档二（s927-45 §4-B：整数字面量＋属性名，不由常量派生）
    assert body_sp.get(qn("w:after")) == DOCX_BODY_AFTER_EFFECTIVE_TWIPS, (
        f"正文段的显式 w:after 现值 {body_sp.get(qn('w:after'))!r} ≠ 字面量 "
        f"{DOCX_BODY_AFTER_EFFECTIVE_TWIPS} —— B′ 接线（docx_exporter.py 两处正文分支）被改动或退回脱钩态；"
        f"若确要改正文段后距，连本断言与 docx_styles.py:24 一起评审"
    )
    assert body_sp.get(qn("w:line")) is not None, (
        "正文段的 w:line 也消失了 —— 读的不是同一个 spacing 元素，本例失去意义"
    )

    defaults_twips = _docdefaults_after_twips(styles_xml)
    # 臂二·等价换轨的哨兵：模板继承值与本仓显式值同格 ⇒ 版式不变（s927-51 §2 实测差异 0 处）
    assert defaults_twips == int(DOCX_BODY_AFTER_EFFECTIVE_TWIPS), (
        f"正文段继承来源（docDefaults）不再是 {DOCX_BODY_AFTER_EFFECTIVE_TWIPS} twips："
        f"实得 {defaults_twips} —— 捆绑模板换了或被 apply_profile 改写过，"
        f"本仓锁定值与模板正式分叉（台账 241ⓒ 记的就是这条未来风险）"
    )
    assert int(body_sp.get(qn("w:after"))) == defaults_twips, (
        f"显式值 {body_sp.get(qn('w:after'))} 与继承值 {defaults_twips} 不同格 ⇒ B′ 的『零版式变化』前提失效，"
        f"正文段观感已被改动（不再是等价换轨）"
    )
    # 臂三·接线因果：产物读数与常量现值同格（容差＝s927-45 §2 的半网格 0.025pt）
    assert int(body_sp.get(qn("w:after"))) / 20.0 == pytest.approx(
        ds.SPACE_AFTER_BODY_PT, abs=0.025
    ), (
        f"正文段显式 w:after 换算 {int(body_sp.get(qn('w:after'))) / 20.0}pt 与常量现值 "
        f"{ds.SPACE_AFTER_BODY_PT}pt 不同格 ⇒ 接线取错了常量/单位（Pt 与 twips 的换算被改动）"
    )

