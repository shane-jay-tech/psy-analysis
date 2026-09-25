# -*- coding: utf-8 -*-
"""`export_to_docx` 的**页面版式 / 字号数值族**xml 属性级快照钉桩（s923-06）。

补的是 `zcode-bridge/done/s922-13-psy-exporters-docx.result.md:95-96` 留下的口子：
「未钉字号/行距/页面设置（`:22-44` 的 `Cm/Pt` 数值族）；那属版式回归，需要图像或 xml 级比对，另单再说」。
本件走的是**xml 属性级比对**（任务书指定），不做图像比对、不引新依赖。

════════ 为什么这条单最容易写出"说谎的断言"（本件的核心口径）════════
源码 `:24-29` 写的是 `Cm(21) / Cm(29.7) / Cm(2.54) / Cm(3.18)` 四个**十进制意图值**，
但 OOXML 的 `w:pgSz` / `w:pgMar` / `w:ind` 存的是 **twips（1/20 磅）整数** ⇒
落盘产物**不可能**逐位等于源码字面量。实测（本机 .venv，python-docx 1.2.0）：

| 源码 | 落盘 twips | 读回 cm | 误差 |
|---|---|---|---|
| `Cm(21)`   | 11906 | 21.000861 | +0.000861 cm |
| `Cm(29.7)` | 16838 | 29.700361 | +0.000361 cm |
| `Cm(3.18)` | 1803  | 3.180292  | +0.000292 cm |
| `Cm(2.54)` | 1440  | **2.54 整** | 0（1 英寸＝整数 twips，唯一免量化的一项）|

⇒ 断言若写成 `section.page_width.cm == 21` **当场是红的**；写成 `== Cm(21)`（Length 比较）也是红的。
本件的口径：**四个十进制值用容差钉（意图面），真实快照用 twips 整数钉（产物面）**，
两层都留，谁漂了都能立刻看出来。**只钉现状，不改版式**（改任何一个数都属版式回归决定）。

同族第三处说谎点曾是 Heading 字号：`font.size` 只写 `w:sz`，模板自带的 `w:szCs`（复杂文种字号）
留在原值 ⇒ 只看 `style.font.size.pt` 会以为 22pt 全线生效。**该条已由 s924-10 一致化**
（`_sync_complex_script_size` 在 Heading 循环里让 `w:szCs` 跟随 `w:sz`），本件第 5 例现在钉的是
"三档两值成对相等"，并**继续钉住样式层的未收口面**：Normal 仍只有 `w:sz`、没有 `szCs`。
另有一处**本件并未断言、只在 result 里登记的实测残留**：封面三个带显式字号的 run 仍无 `w:szCs`
⇒ 那两处要不要一起收，属主班裁（见 s924-10 result 遗留问题）。

红线自约：design 全为「合成构念/合成题干」人造文案；本件**不含任何统计/研究结论数值**，
`scale_config` 里的 5/1/0 只是让导出能跑通的结构参数，不参与任何断言。
"""
from __future__ import annotations

import io
import re
import zipfile

import pytest
from docx import Document
from docx.shared import Cm, Pt

from src.questionnaire import exporters
from src.questionnaire.exporters import export_to_docx

# ── 源码 :24-29 的四个"意图值"（完成标准要求内联的字面量）────────────────────
INTENT_PAGE_WIDTH_CM = 21
INTENT_PAGE_HEIGHT_CM = 29.7
INTENT_MARGIN_TB_CM = 2.54
INTENT_MARGIN_LR_CM = 3.18

# ── 产物真相：twips 整数（OOXML 实际存的）──────────────────────────────────
ACTUAL_PG_W_TWIPS = 11906
ACTUAL_PG_H_TWIPS = 16838
ACTUAL_MARGIN_TB_TWIPS = 1440
ACTUAL_MARGIN_LR_TWIPS = 1803

# 量化误差上界：1 twips = 0.0017639 cm，留 3 twips 余量足够钉住"没人改版式"
TWIPS_GRANULARITY_CM = 0.005


def _design() -> dict:
    return {
        "construct_name": "合成构念",
        "research_question": "合成研究问题？",
        "template_used": {"name": "合成题型"},
        "scale_config": {
            "points": 5, "n_items": 2, "n_dimensions": 1, "n_reverse": 1, "reverse_ratio": "50%",
            "anchors": ["完全不同意", "完全同意"],
        },
        "match_reason": "合成匹配理由",
        "dimensions_used": [{"name": "维度甲", "desc": "合成描述", "item_count": 2}],
        "instructions": "合成指导语",
        "items": [
            {"dimension": "维度甲", "reverse": False, "index": 1, "text": "合成题干一"},
            {"dimension": "维度甲", "reverse": True, "index": 2, "text": "合成题干二"},
        ],
        "scoring": "合成计分方式",
        "psychometrics": {"信度": "合成信度策略", "效度": "合成效度策略"},
    }


@pytest.fixture(scope="module")
def artifact() -> dict:
    """跑一次导出，把三份视图（bytes / 读回 Document / 两个 xml 文本）缓存给整件用。"""
    blob = export_to_docx(_design())
    with zipfile.ZipFile(io.BytesIO(blob)) as zf:
        return {
            "blob": blob,
            "doc": Document(io.BytesIO(blob)),
            "document_xml": zf.read("word/document.xml").decode("utf-8"),
            "styles_xml": zf.read("word/styles.xml").decode("utf-8"),
        }


def _style_xml(styles_xml: str, style_id: str) -> str:
    return re.search(r'<w:style [^>]*w:styleId="%s".*?</w:style>' % style_id, styles_xml, re.S).group(0)


# ═══════ 1. 意图值 vs 读回值：量化这件事本身要钉住 ═══════
def test_page_geometry_readback_matches_intent_only_within_twips_granularity(artifact) -> None:
    """`:24-29` 的四个十进制数**读不回精确值**——只有 2.54（整英寸）能逐位回来。"""
    section = artifact["doc"].sections[0]

    assert section.page_width.cm == pytest.approx(INTENT_PAGE_WIDTH_CM, abs=TWIPS_GRANULARITY_CM)
    assert section.page_height.cm == pytest.approx(INTENT_PAGE_HEIGHT_CM, abs=TWIPS_GRANULARITY_CM)
    assert section.left_margin.cm == pytest.approx(INTENT_MARGIN_LR_CM, abs=TWIPS_GRANULARITY_CM)
    assert section.right_margin.cm == pytest.approx(INTENT_MARGIN_LR_CM, abs=TWIPS_GRANULARITY_CM)
    # 上下边距是四项里唯一"整数友好"的：2.54 cm == 1 英寸 == 1440 twips，量化误差恰好为 0
    assert section.top_margin.cm == INTENT_MARGIN_TB_CM
    assert section.bottom_margin.cm == INTENT_MARGIN_TB_CM

    # 剩下三项的量化方向也钉住（哪天 python-docx 换了舍入策略，这几条会先红）
    assert section.page_width.cm > INTENT_PAGE_WIDTH_CM
    assert section.page_height.cm > INTENT_PAGE_HEIGHT_CM
    assert section.left_margin.cm > INTENT_MARGIN_LR_CM


# ═══════ 2. 权威快照：sectPr 的 twips 整数（任务书要的"xml 属性级比对"）═══════
def test_sectpr_xml_attributes_are_the_authoritative_geometry_snapshot(artifact) -> None:
    """直接读 `word/document.xml` 的 `w:pgSz` / `w:pgMar` 整数属性——**这才是文件真相**。

    用正则而不是只靠 python-docx 的 `.twips`  convenience：属性名与出现顺序也一并钉住，
    版式被人改（例如左右边距从 3.18 挪到 2.54）时这里必红，且红得能指出是哪个属性。
    """
    xml = artifact["document_xml"]
    sect = re.search(r"<w:sectPr\b.*?</w:sectPr>", xml, re.S).group(0)

    assert re.search(r'<w:pgSz w:w="%d" w:h="%d"/>' % (ACTUAL_PG_W_TWIPS, ACTUAL_PG_H_TWIPS), sect), sect
    mar = re.search(
        r'<w:pgMar w:top="(\d+)" w:right="(\d+)" w:bottom="(\d+)" w:left="(\d+)"'
        r' w:header="(\d+)" w:footer="(\d+)" w:gutter="(\d+)"/>', sect)
    assert mar is not None, sect
    top, right, bottom, left, header, footer, gutter = (int(g) for g in mar.groups())
    assert (top, bottom) == (ACTUAL_MARGIN_TB_TWIPS, ACTUAL_MARGIN_TB_TWIPS)      # Cm(2.54)
    assert (left, right) == (ACTUAL_MARGIN_LR_TWIPS, ACTUAL_MARGIN_LR_TWIPS)      # Cm(3.18)
    assert (header, footer, gutter) == (720, 720, 0)
    # 与第 1 例同源自证：twips 换算回 cm 必须落进意图容差
    assert abs(left * 2.54 / 1440 - INTENT_MARGIN_LR_CM) < TWIPS_GRANULARITY_CM


# ═══════ 3. 页眉页脚不是我们设的：别把模板继承当成自己的功劳 ═══════
def test_header_footer_and_gutter_are_inherited_not_set_by_the_exporter(artifact) -> None:
    """`header=720 / footer=720 / gutter=0` **不出自** `:24-29`——那 6 行只设了页宽高与四边距。

    这三项是 python-docx 默认模板带进来的。钉它的目的是防一类误判：
    有人照着"版式已全量钉桩"这句话去改 `:24-29`，会以为页眉距离也在管辖内，其实不在。
    """
    source = open(exporters.__file__, encoding="utf-8").read()
    assert "header_distance" not in source and "footer_distance" not in source and "gutter" not in source

    section = artifact["doc"].sections[0]
    assert section.header_distance.cm == 1.27                                      # 720 twips，模板值
    assert section.footer_distance.cm == 1.27
    assert section.gutter == 0


# ═══════ 4. 字号/行距数值族（半磅单位）═══════
def test_normal_and_heading_size_and_line_spacing_family(artifact) -> None:
    """`:32-41` 的字号族：`w:sz` 存**半磅** ⇒ 12pt 写 24、22/16/14pt 写 44/32/28；行距 1.8 写 432。"""
    styles_xml = artifact["styles_xml"]

    normal = _style_xml(styles_xml, "Normal")
    assert re.search(r'<w:sz w:val="24"/>', normal)                                # Pt(12)
    assert re.search(r'<w:spacing w:line="432" w:lineRule="auto"/>', normal)       # 1.8 × 240
    assert artifact["doc"].styles["Normal"].paragraph_format.line_spacing == 1.8

    for level, half_points in ((1, "44"), (2, "32"), (3, "28")):                   # 22 / 16 / 14 pt
        style = _style_xml(styles_xml, f"Heading{level}")
        assert re.search(r'<w:sz w:val="%s"/>' % half_points, style), level
        assert "<w:b/>" in style                                                   # bold=True
        assert '<w:color w:val="2C3E50"/>' in style                                # RGBColor(0x2C,0x3E,0x50)


# ═══════ 5. 曾经的"常量说谎"实例：sz 跟了、szCs 没跟 —— s924-10 已一致化 ═══════
def test_heading_szcs_follows_sz_after_s92410(artifact) -> None:
    """**已修**（s924-10 ＝ 入册授权）：Heading 1/2/3 的 `w:szCs` 现在与 `w:sz` 同值。

    本例原形是 s923-06 的"只钉不修"现状例。改前实测快照（本机 .venv，值随本断言内联）：
    Heading 1 `sz=44`(22pt) / `szCs=28`(14pt)；Heading 2 `sz=32`(16pt) / `szCs=26`(13pt)；
    Heading 3 `sz=28`(14pt) / **无** `szCs`。⇒ 只看 `style.font.size.pt` 得到 22 会以为全线 22pt，
    复杂文种（阿拉伯文等）其实按 14pt 走 —— 这就是当初钉它的原因。
    现由 `_sync_complex_script_size` 在 Heading 循环里补第三字号槽，三档全跟：44/44、32/32、28/28。

    覆盖面自证（改写后仍有牙，非为绿放宽）：断言的是**成对相等**，不是"存在即可"。
    · 撤掉那次同步 ⇒ H1 回到 `szCs=28`、H3 回到无 `szCs`，三行断言全红；
    · 只同步一档（漏 H3）或写错值 ⇒ 对应那行红。
    **本例同时继续钉住两处未纳入本单的范围**：Normal 仍无 `szCs`、H1 的显式字体名与 theme 属性仍并存
    （后者按卷宗只落 result 裁决注记、不改）。
    """
    styles_xml = artifact["styles_xml"]

    def szs(style_id: str):
        body = _style_xml(styles_xml, style_id)
        return (re.findall(r'<w:sz w:val="(\d+)"/>', body),
                re.findall(r'<w:szCs w:val="(\d+)"/>', body))

    # 完成标准 ①：两值一致，且值就是改前那条 `w:sz`（不是随便相等就算过）
    assert szs("Heading1") == (["44"], ["44"])
    assert szs("Heading2") == (["32"], ["32"])
    assert szs("Heading3") == (["28"], ["28"])                     # 改前该档根本没有 szCs
    assert szs("Normal") == (["24"], [])                           # 本单范围外，仍只钉现状
    # 显式字体名与 theme 属性并存（同为"看着配了"一族）：**本单不动**，仅落 result 裁决注记
    h1 = _style_xml(styles_xml, "Heading1")
    assert 'w:asciiTheme="majorHAnsi"' in h1 and 'w:ascii="黑体"' in h1
    assert 'w:eastAsiaTheme="majorEastAsia"' in h1 and 'w:eastAsia="黑体"' in h1


# ═══════ 6. 正文级版式数值：封面段与缩进的量化 ═══════
def test_cover_paragraph_and_indent_family_numbers(artifact) -> None:
    """`Pt(60)`（:49）与 `Cm(0.74)`（:87/:100/:110/:115/:162）、`Cm(1)`（:148）—— 同样都是"存不下的小数"。

    `w:before` 单位是 1/20 磅 ⇒ Pt(60) = **1200**（整数，无损）；
    `w:firstLine` / `w:left` 单位是 twips ⇒ Cm(0.74) = **420**（读回 0.740833 cm）、
    Cm(1) = **567**（读回 1.000125 cm）。封面主标题 26pt 写 `w:sz="52"`，副标题 16pt 写 `w:sz="32"`。
    """
    xml = artifact["document_xml"]
    doc = artifact["doc"]

    cover = re.findall(r"<w:p\b.*?</w:p>", xml, re.S)[0]
    assert '<w:spacing w:before="1200"/>' in cover               # Pt(60) 无损
    assert '<w:jc w:val="center"/>' in cover
    assert '<w:sz w:val="52"/>' in cover                         # 主标题 Pt(26)
    assert doc.paragraphs[0].paragraph_format.space_before == Pt(60)
    assert doc.paragraphs[0].runs[0].font.size.pt == 26

    subtitle = re.findall(r"<w:p\b.*?</w:p>", xml, re.S)[1]
    assert '<w:sz w:val="32"/>' in subtitle                      # 副标题 Pt(16)
    # s923-05 附带发现（东亚映射缺失）的「两笔一体」改写位，见 s923-19 挡路条目 ＋ s924-09：
    # 生产码补 `_set_east_asian_font(run, "楷体")` 后，本行从 `'w:eastAsia' not in subtitle`
    # 翻成正向钉 —— 撤掉那次补写本例即红（东亚槽读不到 ⇒ 断言落空）。
    assert 'w:eastAsia="楷体"' in subtitle                        # 完成标准 ①：副标题段 XML 含东亚映射

    # 10 = 6 行 info_lines(:87) + 1 段定义(:100) + 1 段维度描述(:110) + 2 道题干(:162)
    assert len(re.findall(r'<w:ind w:firstLine="420"/>', xml)) == 10   # :115 那条本件 design 不触发（无 example）
    assert re.findall(r'<w:ind w:left="(\d+)"/>', xml) == ["567"]        # 指导语 Cm(1)
    indented = [q for q in doc.paragraphs if q.paragraph_format.first_line_indent is not None]
    assert all(q.paragraph_format.first_line_indent.cm == pytest.approx(0.74, abs=TWIPS_GRANULARITY_CM)
               for q in indented) and len(indented) == 10
    assert Cm(1).twips == 567                                       # 换算自证，防"数字抄错"
