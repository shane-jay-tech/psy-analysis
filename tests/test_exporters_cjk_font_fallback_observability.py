# -*- coding: utf-8 -*-
"""`_set_east_asian_font` 的**观测面接线**回归（s924-08，口径 A 落地件）。

对应关系（别和 `tests/test_exporters_cjk_font_fallback.py` 搞混）：
- 那份（s923-05）钉的是「**补齐之前**观测面为零」这个现状，其 docstring 已预告
  「哪天补了 logger，本件 1/2/3 号例会红，届时要把『有 debug 记录』写成新断言而不是删用例」。
- 本件是那份预告所指的那次补齐**之后**的新断言：口径 A（`logger.debug`）已经落地，
  于是钉「每一次真失败都必须留下一条 debug 记录，且回退语义一字未改」。
  两份**同属一笔改动**（生产码 +3 行 ＋ 本件），主班回退时请一起回退。

口径 A 的边界（本件逐条钉住，防止后续人把它长成口径 B/C）：
1. **只记 debug**：级别必须是 DEBUG。抬到 WARNING 会经 Streamlit 侧日志回显到用户界面，
   那是 s923-05 口径 B，明确「需主班裁决、不在执行手授权内」。
2. **不改控制流**：函数仍返回 `None`、仍不抛、仍不 `warnings.warn`。
   ⇒ 这条同时自证「没有把吞掉的异常改成**新的静默路径**」：旧路径吞了不留痕，
      新路径吞了**必留痕**，本件的 records 计数就是痕迹计数器。
3. **不改导出内容**：docx 产物字节长度与 `w:eastAsia` 缺失状态与改前一致。

造失败的手法与那份同型（打 `docx.oxml.ns.qn`，只在被调用时刻生效），
但**刻意 import 复制而非跨文件引用** —— 跨文件依赖会让本件的结论被另一份文件的改动左右，
这个取舍在 s923-04 的 `test_questionnaire_export_pdf_positive.py:49-51` 已立过先例。

红线自约：design 全为「合成构念/合成题干」人造文案，不含真实被试数据、不含统计/研究结论数值。
"""
from __future__ import annotations

import contextlib
import io
import logging
import warnings

import docx.oxml.ns as oxml_ns
from docx import Document
from docx.oxml.ns import qn

from src.questionnaire.exporters import _set_east_asian_font, export_to_docx

_LOGGER_NAME = "src.questionnaire.exporters"
_RPR_TAGS = ("w:rPr", "w:rFonts")
# 一次成功导出里 `_set_east_asian_font` 的**实际调用次数**（不是调用点个数）。
# 与 s923-05 的 5 号例同一口径：Normal 1 ＋ Heading 1/2/3 共 3 ＋ 封面主标题 run 1。
_CALLS_PER_EXPORT = 5


def _design() -> dict:
    return {
        "construct_name": "合成构念",
        "research_question": "合成研究问题？",
        "template_used": {"name": "合成题型"},
        "scale_config": {
            "points": 5, "n_items": 2, "n_dimensions": 1, "n_reverse": 1, "reverse_ratio": "50%",
            "anchors": ["完全不同意", "完全同意"],
        },
        "match_reason": "合成匹配理由",
        "dimensions_used": [{"name": "维度甲", "desc": "合成描述", "item_count": 2}],
        "instructions": "合成指导语",
        "items": [
            {"dimension": "维度甲", "reverse": False, "index": 1, "text": "合成题干一"},
            {"dimension": "维度甲", "reverse": True, "index": 2, "text": "合成题干二"},
        ],
        "scoring": "合成计分方式",
        "psychometrics": {"信度": "合成信度策略", "效度": "合成效度策略"},
    }


@contextlib.contextmanager
def failing_cjk_lookup():
    """把 `docx.oxml.ns.qn` 换成「只对 rPr/rFonts 两个标签抛」⇒ 制造**真**的结构化失败。

    为什么不用「传个垃圾字体名」：`rFonts.set(qn("w:eastAsia"), font_name)` 不校验字体名，
    垃圾名照样塞得进去 ⇒ **根本走不到 except**（s923-05 3 号例已否证过那条口径）。
    命中计数用来区分「真的走到了 except」与「函数压根没被调用」这两种都会表现为零记录的情形。
    """
    real_qn = oxml_ns.qn
    stats = {"hits": 0}

    def _qn(tag: str):
        if tag in _RPR_TAGS:
            stats["hits"] += 1
            raise RuntimeError(f"forced failure on {tag}")
        return real_qn(tag)

    oxml_ns.qn = _qn
    try:
        yield stats
    finally:
        oxml_ns.qn = real_qn


def _records_of(caplog, logger_name: str = _LOGGER_NAME):
    return [r for r in caplog.records if r.name == logger_name]


def _east_asia_of(element):
    """读回 `w:eastAsia`；rPr / rFonts 缺失时分别返回哨兵串（与 s923-05 的 helper 同型）。"""
    rpr = element.find(qn("w:rPr"))
    if rpr is None:
        return "<无 rPr>"
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        return "<无 rFonts>"
    return rfonts.get(qn("w:eastAsia"))


def test_debug_record_written_and_fallback_stays_silent_after_export(caplog) -> None:
    """口径 A 的正题：**每一次真失败都留下一条 DEBUG，且回退仍静默、整篇导出仍成功**。

    断言分三段，缺一都不算接线到位：
    ① 单点失败 ⇒ 恰好 1 条、级别 DEBUG、logger 名与消息载荷（字体名 ＋ 异常文本）都在；
    ② 整篇导出 ⇒ `hits == 5` 且记录数 `== 5` ⇒ **零漏记**（"新的静默路径"若存在，这个等式必破）；
    ③ 导出产物仍可读、正文仍在、`w:eastAsia` 仍缺（失败语义未改），且不抛、不发 warnings。
    """
    # ── ① 单点：留痕 ────────────────────────────────────────────────────────
    style = Document().styles["Normal"]
    with caplog.at_level(logging.DEBUG, logger=_LOGGER_NAME):
        with failing_cjk_lookup() as stats:
            returned = _set_east_asian_font(style, "宋体")

    assert stats["hits"] >= 1, "失败桩没命中 ⇒ 后面所有断言都是空转"
    assert returned is None, "口径 A 不许改返回值（那会扩散到 3 个调用点与 UI 层）"

    recs = _records_of(caplog)
    assert len(recs) == 1, f"应恰好一条 debug，实得 {[(r.name, r.levelname) for r in recs]}"
    assert recs[0].levelno == logging.DEBUG, f"级别被抬了：{recs[0].levelname}"
    assert recs[0].name == _LOGGER_NAME
    msg = recs[0].getMessage()
    assert "宋体" in msg and "forced failure" in msg, msg        # 字体名与异常原因都得可查
    assert "docx 字体" in msg, msg                              # 消息前缀可 grep，供事后按串检索

    caplog.clear()

    # ── ② 全量导出：零漏记 ＋ ③ 产物不受影响 ────────────────────────────────
    with caplog.at_level(logging.DEBUG, logger=_LOGGER_NAME):
        with warnings.catch_warnings(record=True) as caught:
            warnings.simplefilter("always")
            with failing_cjk_lookup() as stats:
                blob = export_to_docx(_design())                 # 不抛＝回退仍静默

    assert stats["hits"] == _CALLS_PER_EXPORT
    recs = _records_of(caplog)
    assert len(recs) == _CALLS_PER_EXPORT, (
        f"记录数 {len(recs)} != 失败次数 {stats['hits']} ⇒ 有失败被吞成新的静默路径"
    )
    assert all(r.levelno == logging.DEBUG for r in recs), [(r.name, r.levelname) for r in recs]
    assert caught == [], "口径 A 不发 warnings：抬到用户可见面属口径 B，需主班裁决"

    doc = Document(io.BytesIO(blob))
    assert len(blob) > 0 and bytes(blob)[:2] == b"PK"            # 仍是能开的 docx
    assert "Q1. 合成题干一" in "\n".join(p.text for p in doc.paragraphs)
    normal = doc.styles["Normal"]
    assert normal.font.name == "宋体"                            # latin 槽照旧（看着"配好了"）
    assert _east_asia_of(normal.element) is None, "失败态下 eastAsia 必须仍写不上（改了它＝改了回退语义）"
