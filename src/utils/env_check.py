"""启动时一键环境检查

毫秒级自检，用 st.toast() 弹出结果，不阻塞启动。
检查项：semopy, factor_analyzer, LLM API Key, 关键依赖。

v5.8 性能优化：所有依赖检查改用 importlib.util.find_spec 探测（只查
安装状态、不真正 import）。此前版本在首次页面加载时同步 import
pingouin/statsmodels/semopy/sklearn 等重型模块（冷启动累计 20s+），
把「5 秒自检」变成了 30 秒白屏。真正的 import 留给功能首次使用时
懒加载完成。
"""

import importlib.metadata
import importlib.util
from typing import Dict, List, Tuple


def _probe(package: str, import_name: str | None = None) -> Tuple[bool, str]:
    """探测包是否安装（find_spec 不执行模块代码，毫秒级）。

    Args:
        package: PyPI 包名（用于显示与版本查询）。
        import_name: 顶层 import 名，缺省与 package 相同。

    Returns:
        (ok, message)。ok=False 时 message 为「未安装」提示。
    """
    spec = importlib.util.find_spec(import_name or package)
    if spec is None:
        return False, f"{package} 未安装"
    try:
        ver = importlib.metadata.version(package)
        return True, f"{package} {ver}"
    except importlib.metadata.PackageNotFoundError:
        return True, f"{package} (已安装)"


def check_semopy() -> Tuple[bool, str]:
    """检查 semopy (CFA/SEM) 是否可用（仅探测，不 import）"""
    ok, msg = _probe("semopy")
    if not ok:
        return False, "semopy 未安装 — CFA 功能不可用，EFA 正常"
    return ok, msg


def check_kaleido() -> Tuple[bool, str]:
    """检查 kaleido (Plotly 图表静态导出) 是否可用（仅探测，不 import）"""
    ok, msg = _probe("kaleido")
    if not ok:
        return False, "kaleido 未安装 — 论文版 PNG 导出不可用，HTML 导出正常"
    return ok, msg


def check_factor_analyzer() -> Tuple[bool, str]:
    """检查 factor_analyzer 版本（仅探测，不 import）"""
    ok, msg = _probe("factor-analyzer", import_name="factor_analyzer")
    if not ok:
        return False, "factor_analyzer 未安装 — EFA 功能不可用"
    return ok, msg


def check_llm_api() -> Tuple[bool, str]:
    """检查 D:\\code\\.env.local 是否至少配了一个快速模型。"""
    try:
        from src.llm_gateway.quick_models import list_available_quick_models
        models = list_available_quick_models()
        available = [m for m in models if m.get("available")]
        if available:
            names = "、".join(m["label"] for m in available[:2])
            extra = f"…等 {len(available)} 个" if len(available) > 2 else ""
            return True, f"快速模型已配置：{names}{extra}"
        return False, r"D:\code\.env.local 中未配置任何模型"
    except Exception as e:
        return False, f"快速模型检查失败: {e}"


def check_critical_deps() -> List[Tuple[str, bool, str]]:
    """检查关键依赖（find_spec 探测，不 import，毫秒级）"""
    deps = [
        ("pandas", "pandas"),
        ("numpy", "numpy"),
        ("scipy", "scipy"),
        ("statsmodels", "statsmodels"),
        ("plotly", "plotly"),
        ("streamlit", "streamlit"),
        ("pingouin", "pingouin"),
        ("openpyxl", "openpyxl"),
        ("scikit-learn", "sklearn"),
    ]
    results = []
    for name, import_name in deps:
        spec = importlib.util.find_spec(import_name)
        if spec is None:
            results.append((name, False, "未安装"))
            continue
        try:
            ver = importlib.metadata.version(name if name != "scikit-learn" else "scikit-learn")
            results.append((name, True, ver))
        except importlib.metadata.PackageNotFoundError:
            results.append((name, True, "已安装"))
    return results


def run_startup_check() -> Dict:
    """
    运行完整启动检查，返回状态字典。

    {"semopy_ok": bool, "factor_analyzer_ok": bool, "llm_ok": bool,
     "critical_ok": bool, "warnings": [...], "errors": [...]}
    """
    status = {
        "semopy_ok": False,
        "factor_analyzer_ok": False,
        "kaleido_ok": False,
        "llm_ok": False,
        "critical_ok": True,
        "warnings": [],
        "errors": [],
        "details": {},
    }

    # semopy
    ok, msg = check_semopy()
    status["semopy_ok"] = ok
    status["details"]["semopy"] = msg
    if not ok:
        status["warnings"].append("⚠️ CFA/SEM 不可用（semopy 缺失），请运行：pip install semopy")

    # kaleido (论文版 PNG 导出)
    ok, msg = check_kaleido()
    status["kaleido_ok"] = ok
    status["details"]["kaleido"] = msg
    if not ok:
        status["warnings"].append("⚠️ 论文版 PNG 导出不可用（kaleido 缺失），请运行：pip install kaleido")

    # factor_analyzer
    ok, msg = check_factor_analyzer()
    status["factor_analyzer_ok"] = ok
    status["details"]["factor_analyzer"] = msg
    if not ok:
        status["warnings"].append("⚠️ EFA 不可用（factor_analyzer 缺失）")

    # LLM
    ok, msg = check_llm_api()
    status["llm_ok"] = ok
    status["details"]["llm"] = msg
    if not ok:
        status["warnings"].append(
            r"💡 请在 D:\code\.env.local 配置模型（GPT_/DEEPSEEK_/KIMI_/CLAUDE_ "
            "三件套），然后顶部「🤖 AI 模型」选一个。模板：D:\\code\\.env.local.example"
        )

    # 关键依赖
    dep_results = check_critical_deps()
    for name, ok, ver in dep_results:
        status["details"][name] = ver
        if not ok:
            status["critical_ok"] = False
            status["errors"].append(f"❌ {name} 未安装 — 核心功能可能异常")

    return status


def render_env_status_toasts(status: Dict):
    """用 st.toast() 显示启动检查结果"""
    import streamlit as st
    import time

    if status["errors"]:
        for err in status["errors"]:
            st.toast(err, icon="❌")
        time.sleep(0.3)

    if status["warnings"]:
        for w in status["warnings"]:
            st.toast(w, icon="⚠️")
        time.sleep(0.3)

    if not status["errors"] and not status["warnings"]:
        st.toast("✅ 环境检查通过，所有核心功能可用", icon="✅")
    elif not status["errors"]:
        st.toast("⚠️ 部分可选功能不可用，核心功能正常", icon="⚠️")
    else:
        st.toast("❌ 关键依赖缺失，请检查环境", icon="❌")


# --------------------------------------------------------------------------- #
# v2.9: 深度自检（生成测试 PDF/Word 验证可用性）
# --------------------------------------------------------------------------- #

def deep_check_pdf_generation() -> Tuple[bool, str]:
    """实际生成一个最小 PDF，验证 fpdf2 + CJK 字体能正常工作。"""
    try:
        from fpdf import FPDF
        cjk_path = _find_cjk_font_for_check()
        pdf = FPDF()
        pdf.add_page()
        if cjk_path:
            pdf.add_font("CJK", "", cjk_path)
            pdf.set_font("CJK", "", 12)
            pdf.cell(0, 10, "中文字体测试")
        else:
            pdf.set_font("Helvetica", "", 12)
            pdf.cell(0, 10, "PDF generation OK (no CJK font)")
        out = pdf.output()
        if isinstance(out, str):
            out = out.encode("latin-1")
        if out and out[:4] == b"%PDF":
            return True, ("PDF 生成正常" + ("（含中文字体）" if cjk_path else "（无中文字体）"))
        return False, "PDF 生成失败：返回数据非合法 PDF"
    except Exception as e:
        return False, f"PDF 生成异常：{e}"


def deep_check_docx_generation() -> Tuple[bool, str]:
    """实际生成一个最小 docx，验证 python-docx 能正常工作。"""
    try:
        from docx import Document
        import io
        doc = Document()
        doc.add_paragraph("Word generation test - 中文测试")
        buf = io.BytesIO()
        doc.save(buf)
        data = buf.getvalue()
        if data[:2] == b"PK":
            return True, "Word 生成正常"
        return False, "Word 生成失败：返回数据非合法 docx"
    except Exception as e:
        return False, f"Word 生成异常：{e}"


def _find_cjk_font_for_check() -> str:
    """简化版 CJK 字体查找，仅用于 env_check。"""
    import os
    win_dir = os.environ.get("WINDIR", "C:\\Windows")
    candidates = [
        os.path.join(win_dir, "Fonts", "msyh.ttc"),
        os.path.join(win_dir, "Fonts", "simhei.ttf"),
        "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
        "/System/Library/Fonts/PingFang.ttc",
    ]
    for p in candidates:
        if os.path.exists(p):
            return p
    return ""


def run_deep_environment_check(fast: bool = True) -> Dict:
    """v2.9: 深度环境自检，返回详细的环境健康报告。

    用于 UI 顶部橙色提示条：任一项失败即提示用户。

    v5.8: fast=True（默认，启动提示条路径）时跳过「生成测试 PDF/Word」
    等秒级实测，仅做毫秒级探测；完整实测只在用户主动点「运行系统诊断」
    时执行（fast=False），避免每次会话启动都被 PDF/Word 生成阻塞。
    """
    report = {
        "all_ok": True,
        "checks": [],  # list of (name, ok, message)
        "fix_actions": [],  # list of "建议运行：pip install xxx"
    }

    # 基础依赖
    sem_ok, sem_msg = check_semopy()
    report["checks"].append(("semopy (CFA)", sem_ok, sem_msg))
    if not sem_ok:
        report["fix_actions"].append("pip install -U semopy  # 解锁 CFA 功能")

    fa_ok, fa_msg = check_factor_analyzer()
    report["checks"].append(("factor_analyzer (EFA)", fa_ok, fa_msg))
    if not fa_ok:
        report["all_ok"] = False  # 必备
        report["fix_actions"].append("pip install -U factor-analyzer  # 必装")

    kal_ok, kal_msg = check_kaleido()
    report["checks"].append(("kaleido (PNG 导出)", kal_ok, kal_msg))
    if not kal_ok:
        report["fix_actions"].append("pip install -U kaleido  # 解锁论文版图表导出")

    # CJK 字体
    cjk_path = _find_cjk_font_for_check()
    cjk_ok = bool(cjk_path)
    report["checks"].append((
        "CJK 字体（Word/PDF 中文）",
        cjk_ok,
        f"已找到 {cjk_path}" if cjk_ok else "未找到中文字体（PDF 将降级为英文）",
    ))
    if not cjk_ok:
        report["fix_actions"].append("Linux 用户运行：sudo apt install fonts-noto-cjk")

    # 深度检查：生成测试 PDF/Word（仅完整诊断时实测）
    if not fast:
        pdf_ok, pdf_msg = deep_check_pdf_generation()
        report["checks"].append(("PDF 生成", pdf_ok, pdf_msg))
        if not pdf_ok:
            report["all_ok"] = False
            report["fix_actions"].append("pip install -U fpdf2")

        docx_ok, docx_msg = deep_check_docx_generation()
        report["checks"].append(("Word 生成", docx_ok, docx_msg))
        if not docx_ok:
            report["all_ok"] = False
            report["fix_actions"].append("pip install -U python-docx")

    return report


def render_env_health_banner():
    """v2.9: 在 UI 顶部显示橙色环境健康提示条（仅在有问题时）。"""
    import streamlit as st

    # 缓存到 session_state，避免每次 rerun 都跑深度检查
    if "_env_deep_check" not in st.session_state:
        st.session_state["_env_deep_check"] = run_deep_environment_check()

    report = st.session_state["_env_deep_check"]

    failures = [c for c in report["checks"] if not c[1]]
    if not failures:
        return  # 一切正常，不显示

    summary = "、".join(name for name, _, _ in failures[:3])
    if len(failures) > 3:
        summary += f"…等 {len(failures)} 项"

    with st.container():
        st.warning(
            f"⚠️ 检测到环境问题：**{summary}**，部分功能可能受影响。"
            "点击下方查看具体解决方案。"
        )
        with st.expander("🔧 查看环境检查详情与修复建议", expanded=False):
            for name, ok, msg in report["checks"]:
                icon = "✅" if ok else "❌"
                st.text(f"{icon} {name}：{msg}")
            if report["fix_actions"]:
                st.markdown("**建议执行：**")
                for action in report["fix_actions"]:
                    st.code(action, language="bash")
